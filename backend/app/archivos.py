"""Generación de reportes descargables: PDF (reportlab) y Excel (openpyxl)."""
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.graphics.shapes import Drawing, Rect, String
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

# Paleta de los mockups
CAFE = "4A3426"
VERDE = "6F8466"
CREMA = "F5F1E8"
ROJO = "B5483B"
AMBAR = "C98A2E"

MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
         "septiembre", "octubre", "noviembre", "diciembre"]


def nombre_periodo(periodo: str) -> str:
    anio, mes = periodo.split("-")
    return f"{MESES[int(mes) - 1].capitalize()} {anio}"


def pesos(valor) -> str:
    return "$" + f"{int(round(float(valor or 0))):,}".replace(",", ".")


def numero(valor) -> str:
    v = float(valor or 0)
    texto = f"{v:,.2f}".rstrip("0").rstrip(".") if v != int(v) else f"{int(v):,}"
    return texto.replace(",", "X").replace(".", ",").replace("X", ".")


# =====================================================================
# PDF: reporte mensual
# =====================================================================
def _barras_categoria(categorias: list[dict], ancho: float) -> Drawing:
    alto_fila, etiqueta, valor_w = 18, 3.2 * cm, 2.8 * cm
    d = Drawing(ancho, alto_fila * max(len(categorias), 1) + 4)
    maximo = max((float(c["consumo"]) for c in categorias), default=0) or 1
    barra_max = ancho - etiqueta - valor_w
    for i, c in enumerate(categorias):
        y = d.height - (i + 1) * alto_fila
        d.add(String(0, y + 5, c["categoria"], fontName="Helvetica", fontSize=9))
        d.add(Rect(etiqueta, y + 3, barra_max, 10, fillColor=colors.HexColor("#EDE6D8"), strokeColor=None))
        d.add(Rect(etiqueta, y + 3, barra_max * float(c["consumo"]) / maximo, 10,
                   fillColor=colors.HexColor("#" + VERDE), strokeColor=None))
        d.add(String(ancho, y + 5, pesos(c["consumo"]), fontName="Helvetica-Bold", fontSize=9, textAnchor="end"))
    return d


def _tabla(filas: list[list], anchos: list[float], alinear_derecha: tuple[int, ...] = ()) -> Table:
    t = Table(filas, colWidths=anchos, repeatRows=1)
    estilo = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + CAFE)),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#" + CREMA)]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D8CFBF")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    for col in alinear_derecha:
        estilo.append(("ALIGN", (col, 0), (col, -1), "RIGHT"))
    t.setStyle(TableStyle(estilo))
    return t


def pdf_reporte_mensual(datos: dict) -> bytes:
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=1.8 * cm, rightMargin=1.8 * cm,
                            topMargin=1.6 * cm, bottomMargin=1.6 * cm,
                            title=f"Reporte mensual {datos['periodo']}", author="Taller Irribarra")
    ancho = A4[0] - doc.leftMargin - doc.rightMargin
    estilos = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=estilos["Title"], alignment=TA_LEFT, fontSize=18,
                        textColor=colors.HexColor("#" + CAFE), spaceAfter=2)
    sub = ParagraphStyle("sub", parent=estilos["Normal"], fontSize=9, textColor=colors.HexColor("#6B5E50"))
    h2 = ParagraphStyle("h2", parent=estilos["Heading2"], fontSize=11.5, spaceBefore=12, spaceAfter=6,
                        textColor=colors.HexColor("#" + CAFE))
    normal = ParagraphStyle("n", parent=estilos["Normal"], fontSize=8.5)

    r = datos["resumen"]
    historia = [
        Paragraph("Taller Irribarra · Saunas · Tinas · Wellness", sub),
        Paragraph(f"Reporte mensual — {nombre_periodo(datos['periodo'])}", h1),
        Paragraph(f"Generado el {datetime.now():%d-%m-%Y %H:%M} por {datos['generado_por']}", sub),
        Spacer(1, 10),
    ]

    # Indicadores principales
    kpis = [
        ("CONSUMO DE MATERIALES", pesos(r["consumo_total_materiales"])),
        ("ALERTAS GENERADAS", str(r["total_alertas"])),
        ("AHORRO POR RECICLAJE", pesos(r["ahorro_por_reciclaje"])),
        ("COMANDAS FINALIZADAS", str(r["comandas_finalizadas"])),
    ]
    celdas = [[Paragraph(f'<font size="7" color="#6B5E50">{t}</font><br/><br/>'
                         f'<font size="15"><b>{v}</b></font>', normal) for t, v in kpis]]
    tk = Table(celdas, colWidths=[ancho / 4] * 4, rowHeights=[1.7 * cm])
    tk.setStyle(TableStyle([
        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#D8CFBF")),
        ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D8CFBF")),
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#FBF9F4")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    historia.append(tk)

    historia.append(Paragraph("Consumo de materiales por categoría", h2))
    if datos["por_categoria"]:
        historia.append(_barras_categoria(datos["por_categoria"], ancho))
    else:
        historia.append(Paragraph("Sin consumo registrado en el período.", normal))

    historia.append(Paragraph("Materiales más consumidos", h2))
    if datos["top_materiales"]:
        filas = [["Código", "Material", "Cantidad", "Unidad", "Costo"]] + [
            [m["codigo_interno"], m["material"], numero(m["cantidad"]), m["unidad_medida"], pesos(m["costo"])]
            for m in datos["top_materiales"]]
        historia.append(_tabla(filas, [2.6 * cm, ancho - 9.6 * cm, 2.2 * cm, 1.8 * cm, 3 * cm], (2, 4)))
    else:
        historia.append(Paragraph("Sin salidas por fabricación en el período.", normal))

    historia.append(Paragraph("Alertas de stock del período", h2))
    if datos["alertas"]:
        filas = [["Fecha", "Tipo", "Material", "Disponible", "Requerido", "Nota", "Estado"]] + [
            [f"{a['fecha_alerta']:%d-%m}", a["tipo"], a["material"], numero(a["stock_disponible"]),
             numero(a["stock_requerido"]), a["folio"] or "—", a["estado"]] for a in datos["alertas"]]
        historia.append(_tabla(filas, [1.4 * cm, 2.2 * cm, ancho - 12.4 * cm, 2 * cm, 2 * cm, 1.6 * cm, 3.2 * cm],
                               (3, 4)))
    else:
        historia.append(Paragraph("No se generaron alertas.", normal))

    historia.append(Paragraph("Mermas reutilizadas (ahorro)", h2))
    if datos["mermas"]:
        filas = [["Retazo", "Material", "Medidas (cm)", "Comanda", "Ahorro"]] + [
            [f"#{m['id_merma']:04d}", m["material"],
             f"{numero(m['largo_cm'])} × {numero(m['ancho_cm'])} × {numero(m['espesor_cm'])}",
             str(m["id_comanda_uso"] or "—"), pesos(m["valor_estimado"])] for m in datos["mermas"]]
        historia.append(_tabla(filas, [1.8 * cm, ancho - 10.8 * cm, 3.6 * cm, 2.2 * cm, 3.2 * cm], (4,)))
    else:
        historia.append(Paragraph("No se reutilizaron retazos en el período.", normal))

    def pie(canvas, documento):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#6B5E50"))
        canvas.drawString(documento.leftMargin, 1 * cm, "Sistema de Gestión de Inventario · Taller Irribarra")
        canvas.drawRightString(A4[0] - documento.rightMargin, 1 * cm, f"Página {documento.page}")
        canvas.restoreState()

    doc.build(historia, onFirstPage=pie, onLaterPages=pie)
    return buffer.getvalue()


# =====================================================================
# Excel
# =====================================================================
_BORDE = Border(bottom=Side(style="thin", color="D8CFBF"))
_ENCABEZADO = PatternFill("solid", fgColor=CAFE)
_RELLENO_ESTADO = {
    "OK": PatternFill("solid", fgColor="E3EEDC"),
    "Bajo": PatternFill("solid", fgColor="FBEFD5"),
    "Critico": PatternFill("solid", fgColor="F8DCD6"),
    "Sin stock": PatternFill("solid", fgColor="F2C4BC"),
}


def _titulo(ws, titulo: str, subtitulo: str, columnas: int) -> None:
    ws["A1"] = titulo
    ws["A1"].font = Font(size=14, bold=True, color=CAFE)
    ws["A2"] = subtitulo
    ws["A2"].font = Font(size=9, color="6B5E50")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columnas)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=columnas)


def _encabezados(ws, fila: int, nombres: list[str], anchos: list[int]) -> None:
    for i, (nombre, ancho) in enumerate(zip(nombres, anchos), start=1):
        c = ws.cell(row=fila, column=i, value=nombre)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = _ENCABEZADO
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.row_dimensions[fila].height = 28
    ws.freeze_panes = ws.cell(row=fila + 1, column=1)


def _a_bytes(wb: Workbook) -> bytes:
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def excel_inventario(materiales: list[dict], alertas: list[dict], usuario: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventario"
    cols = ["Código", "Material", "Tipo", "Bodega", "Unidad", "Stock", "Comprometido", "Disponible",
            "Mínimo", "Costo unitario", "Valor stock", "Estado"]
    _titulo(ws, "Inventario de materiales e insumos — Taller Irribarra",
            f"Generado el {datetime.now():%d-%m-%Y %H:%M} por {usuario}", len(cols))
    _encabezados(ws, 4, cols, [14, 34, 13, 19, 8, 10, 13, 11, 10, 13, 14, 11])

    fila = 5
    for m in materiales:
        valores = [m["codigo_interno"], m["nombre"], m["tipo"], m["bodega"], m["unidad_medida"],
                   float(m["stock_actual"]), float(m["comprometido"]), float(m["disponible"]),
                   float(m["stock_minimo"]), float(m["costo_unitario"]), None, m["estado_stock"]]
        for col, v in enumerate(valores, start=1):
            c = ws.cell(row=fila, column=col, value=v)
            c.border = _BORDE
        ws.cell(row=fila, column=11, value=f"=F{fila}*J{fila}")
        for col in (6, 7, 8, 9):
            ws.cell(row=fila, column=col).number_format = "#,##0.##"
        for col in (10, 11):
            ws.cell(row=fila, column=col).number_format = '"$"#,##0'
        ws.cell(row=fila, column=12).fill = _RELLENO_ESTADO.get(m["estado_stock"], PatternFill())
        fila += 1

    ultima = fila - 1
    ws.cell(row=fila, column=10, value="Total").font = Font(bold=True)
    total = ws.cell(row=fila, column=11, value=f"=SUM(K5:K{ultima})")
    total.font = Font(bold=True)
    total.number_format = '"$"#,##0'
    ws.auto_filter.ref = f"A4:L{ultima}"

    # Hoja de alertas pendientes
    wa = wb.create_sheet("Alertas pendientes")
    cols_a = ["Fecha", "Tipo", "Código", "Material", "Disponible", "Requerido", "Faltante", "Unidad", "Nota de venta"]
    _titulo(wa, "Alertas de stock pendientes", f"Total: {len(alertas)}", len(cols_a))
    _encabezados(wa, 4, cols_a, [17, 12, 14, 34, 11, 11, 11, 8, 13])
    for i, a in enumerate(alertas, start=5):
        valores = [a["fecha_alerta"], a["tipo"], a["codigo_interno"], a["material"], float(a["stock_disponible"]),
                   float(a["stock_requerido"]), float(a["faltante"]), a["unidad_medida"], a["folio"] or "—"]
        for col, v in enumerate(valores, start=1):
            c = wa.cell(row=i, column=col, value=v)
            c.border = _BORDE
        wa.cell(row=i, column=1).number_format = "dd-mm-yyyy hh:mm"

    # Resumen por bodega y estado
    wr = wb.create_sheet("Resumen")
    _titulo(wr, "Resumen por bodega", "Valorizado a costo unitario actual", 4)
    _encabezados(wr, 4, ["Bodega", "Materiales", "Bajo mínimo", "Valor stock"], [22, 12, 13, 16])
    resumen: dict[str, list] = {}
    for m in materiales:
        r = resumen.setdefault(m["bodega"], [0, 0, 0.0])
        r[0] += 1
        r[1] += m["estado_stock"] != "OK"
        r[2] += float(m["valor_stock"])
    for i, (bodega, (n, bajos, valor)) in enumerate(resumen.items(), start=5):
        wr.cell(row=i, column=1, value=bodega)
        wr.cell(row=i, column=2, value=n)
        wr.cell(row=i, column=3, value=bajos)
        wr.cell(row=i, column=4, value=valor).number_format = '"$"#,##0'
    return _a_bytes(wb)


def excel_kardex(material: dict, movimientos: list[dict], usuario: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kardex"
    cols = ["Fecha", "Tipo de movimiento", "Referencia", "Entrada", "Salida", "Saldo", "Costo unitario", "Usuario"]
    _titulo(ws, f"Kardex — {material['codigo_interno']} {material['nombre']}",
            f"Unidad: {material['unidad_medida']} · Bodega: {material['bodega']} · "
            f"Generado el {datetime.now():%d-%m-%Y %H:%M} por {usuario}", len(cols))
    _encabezados(ws, 4, cols, [17, 20, 22, 11, 11, 11, 14, 14])
    nombres_tipo = {"CARGA_INICIAL": "Carga inicial", "ENTRADA_FACTURA": "Entrada por factura",
                    "SALIDA_FABRICACION": "Salida a fabricación", "AJUSTE": "Ajuste de inventario"}
    for i, mv in enumerate(movimientos, start=5):
        cantidad = float(mv["cantidad"])
        valores = [mv["fecha"], nombres_tipo.get(mv["tipo"], mv["tipo"]), mv["referencia"],
                   cantidad if cantidad > 0 else None, -cantidad if cantidad < 0 else None,
                   float(mv["stock_resultante"]), float(mv["costo_unitario"]), mv["usuario"] or "sistema"]
        for col, v in enumerate(valores, start=1):
            c = ws.cell(row=i, column=col, value=v)
            c.border = _BORDE
        ws.cell(row=i, column=1).number_format = "dd-mm-yyyy hh:mm"
        for col in (4, 5, 6):
            ws.cell(row=i, column=col).number_format = "#,##0.##"
        ws.cell(row=i, column=7).number_format = '"$"#,##0'
    return _a_bytes(wb)
